import customtkinter as ctk
from tkinter import filedialog
import os
import re
import glob
import pandas as pd
import fitz
import threading
import time
import openpyxl
from pypdf import PdfReader

# Настройки внешнего вида
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("green")

# --- ФУНКЦИИ ПЕРВОЙ ВКЛАДКИ (QR-СКАНЕР С ОБЩИМ PDF) ---
def process_all_pdfs(pdf_dir, output_dir, qrs_to_find, log_callback, progress_callback):
    pdf_files = glob.glob(os.path.join(pdf_dir, '*.pdf'))
    pdf_files = [f for f in pdf_files if not os.path.basename(f).startswith('Matched_') and not os.path.basename(f).startswith('All_Matched_')]
    
    total_matched_pages = 0
    processed_files_count = 0
    found_qrs = set() 
    
    # Создаем общий документ для объединения всех найденных страниц
    master_doc = fitz.open()
    
    for i, pdf_path in enumerate(pdf_files):
        filename = os.path.basename(pdf_path)
        progress_callback(i + 1, len(pdf_files), filename)
        
        try:
            doc = fitz.open(pdf_path)
            new_doc = fitz.open()
            file_matched_count = 0
            
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                if text:
                    clean_text = re.sub(r'\s+', '', text)
                    for qr in qrs_to_find:
                        clean_qr = re.sub(r'\s+', '', qr)
                        if clean_qr in clean_text:
                            # Добавляем страницу в индивидуальный файл
                            new_doc.insert_pdf(doc, from_page=page_num, to_page=page_num)
                            # И сразу добавляем эту же страницу в общий мастер-документ
                            master_doc.insert_pdf(doc, from_page=page_num, to_page=page_num)
                            
                            file_matched_count += 1
                            total_matched_pages += 1
                            found_qrs.add(qr) 
                            break
                            
            if file_matched_count > 0:
                output_filename = f'Matched_{filename}'
                output_path = os.path.join(output_dir, output_filename)
                new_doc.save(output_path)
                log_callback(f"Успех! Найдено {file_matched_count} стр. -> Сохранен как: {output_filename}")
                processed_files_count += 1
                
            new_doc.close()
            doc.close()
                
        except Exception as e:
            log_callback(f"Ошибка при чтении файла {filename}: {e}")
            
    # Сохраняем общий сводный PDF, если набрались страницы
    if total_matched_pages > 0:
        all_pdf_path = os.path.join(output_dir, "All_Matched_Report.pdf")
        master_doc.save(all_pdf_path)
        log_callback(f"\n[ИТОГ] Создан общий файл со всеми найденными страницами: All_Matched_Report.pdf")
        
    master_doc.close()
    return total_matched_pages, processed_files_count, found_qrs


# --- ФУНКЦИИ ВТОРОЙ ВКЛАДКИ (ПАРСЕР OZON PDF) ---
def parse_orders_from_pdf(pdf_path: str) -> pd.DataFrame:
    """Извлекает текст из PDF и корректно парсит любые форматы артикулов, отправления и товары Ozon."""
    reader = PdfReader(pdf_path)
    full_text = ""
    for page in reader.pages:
        text = page.extract_text()
        if text:
            full_text += text + "\n"

    lines = full_text.split('\n')
    skus = []
    orders = []
    products = []

    for i, line in enumerate(lines):
        line_clean = line
        
        # 1. СНАЧАЛА ИЩЕМ НОМЕР ОТПРАВЛЕНИЯ (и удаляем его из строки, чтобы не путался с артикулом)
        m_order = re.search(r'\d{7,11}-\d{4}-\d{1,2}', line)
        if m_order:
            orders.append((i, m_order.group(0)))
            line_clean = line_clean.replace(m_order.group(0), '')
            
        # 2. УНИВЕРСАЛЬНЫЙ ПОИСК АРТИКУЛА (в очищенной строке)
        m_sku = re.search(r'(\b\d+[_-]\d+\b|\b\d{5,}\b|\b\d{3}[ _]\d{2}(?:\s*\([^)]+\))?\b)', line_clean)
        if m_sku:
            sku_val = m_sku.group(1).replace(' ', '_')
            skus.append((i, sku_val))
            
        # 3. ПОИСК И ОЧИСТКА НАЗВАНИЯ ТОВАРА
        prod = line_clean.replace('|', '').strip()
        if m_sku:
            prod = prod.replace(m_sku.group(0), '').strip()
            
        # Убираем хвосты Ozon (Кол-во + Этикетка, например " 1 7116" в конце строки)
        prod = re.sub(r'\s+\d+\s+\d{4}$', '', prod).strip()
        # Убираем технический мусор (например, "1 -1" или "2 -2" в начале строки)
        prod = re.sub(r'^[-_\s\d]+', '', prod).strip()
        
        # Если после всех чисток осталось осмысленное название (длиннее 2 букв)
        if prod and len(prod) > 2 and not prod.replace('.', '', 1).isdigit():
            if not re.match(r'^(шт|руб|\d+\s*шт|\d+\s*руб)$', prod, re.IGNORECASE):
                products.append((i, prod))

    # Связываем найденные артикулы с ближайшими заказами и товарами по индексу строки
    records = []
    for sku_idx, sku_val in skus:
        closest_order = min(orders, key=lambda x: abs(x[0] - sku_idx)) if orders else (0, '')
        closest_product = min(products, key=lambda x: abs(x[0] - sku_idx)) if products else (0, 'Товар не указан')
        
        order_val = closest_order[1]
        label = order_val.split('-')[0][-4:] if '-' in order_val else ''
        
        records.append({
            'Номер отправления': order_val,
            'Товар': closest_product[1],
            'Артикул': sku_val,
            'Кол-во': 1,
            'Этикетка': label
        })

    df = pd.DataFrame(records)
    if not df.empty:
        df.index = df.index + 1
        df.index.name = '№'
    return df

def save_to_formatted_excel(df: pd.DataFrame, output_path: str):
    """Сохраняет DataFrame в красивый Excel-файл с разметкой."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказы"

    headers = ['№', 'Номер отправления', 'Товар', 'Артикул', 'Кол-во', 'Этикетка']
    ws.append(headers)

    for row in df.itertuples():
        ws.append([row.Index, row[1], row[2], row[3], row[4], row[5]])

    header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="3B5998", end_color="3B5998", fill_type="solid")
    thin_border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin', color="DDDDDD"), 
        right=openpyxl.styles.Side(style='thin', color="DDDDDD"), 
        top=openpyxl.styles.Side(style='thin', color="DDDDDD"), 
        bottom=openpyxl.styles.Side(style='thin', color="DDDDDD")
    )

    for col in range(1, 7):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 12

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
            cell.alignment = openpyxl.styles.Alignment(vertical="center", horizontal="left" if cell.column == 3 else "center")
            if cell.row % 2 == 0:
                cell.fill = openpyxl.styles.PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

    wb.save(output_path)


# --- ГЛАВНОЕ ОКНО ПРИЛОЖЕНИЯ С ВКЛАДКАМИ ---
class MainApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Инструменты логистики")
        self.geometry("850x850")

        self.tabview = ctk.CTkTabview(self, width=810, height=810)
        self.tabview.pack(padx=20, pady=20, fill="both", expand=True)

        self.tab_qr = self.tabview.add("Сверка QR-кодов")
        self.tab_ozon = self.tabview.add("Парсер Ozon PDF")

        self.init_qr_tab()
        self.init_ozon_tab()

    # --- ИНТЕРФЕЙС ВКЛАДКИ 1 (QR) ---
    def init_qr_tab(self):
        header = ctk.CTkLabel(self.tab_qr, text="Сверка QR-кодов по PDF", font=ctk.CTkFont(size=22, weight="bold"), text_color="#90CAF9")
        header.pack(pady=(10, 15))

        files_frame = ctk.CTkFrame(self.tab_qr, fg_color="transparent")
        files_frame.pack(fill="x", padx=10, pady=5)

        self.btn_excel = ctk.CTkButton(files_frame, text="Выбрать Excel", width=150, command=self.pick_excel)
        self.btn_excel.grid(row=0, column=0, padx=(0, 10), pady=10)
        self.entry_excel = ctk.CTkEntry(files_frame, placeholder_text="Файл Excel не выбран", width=570)
        self.entry_excel.grid(row=0, column=1, pady=10)

        self.btn_pdf_in = ctk.CTkButton(files_frame, text="Исходные PDF", width=150, command=self.pick_pdf_in)
        self.btn_pdf_in.grid(row=1, column=0, padx=(0, 10), pady=10)
        self.entry_pdf_in = ctk.CTkEntry(files_frame, placeholder_text="Папка с исходными PDF не выбрана", width=570)
        self.entry_pdf_in.grid(row=1, column=1, pady=10)

        self.btn_pdf_out = ctk.CTkButton(files_frame, text="Куда сохранить?", width=150, command=self.pick_pdf_out)
        self.btn_pdf_out.grid(row=2, column=0, padx=(0, 10), pady=10)
        self.entry_pdf_out = ctk.CTkEntry(files_frame, placeholder_text="Папка для сохранения не выбрана", width=570)
        self.entry_pdf_out.grid(row=2, column=1, pady=10)

        self.btn_start = ctk.CTkButton(self.tab_qr, text="Запустить сверку", height=42, width=250, font=ctk.CTkFont(size=15), command=self.start_qr_processing)
        self.btn_start.pack(pady=15)

        self.lbl_stage = ctk.CTkLabel(self.tab_qr, text="Стадия: Готово к работе", text_color="#ffcc00", font=ctk.CTkFont(weight="bold"))
        self.lbl_stage.pack(pady=5)

        self.progress_bar = ctk.CTkProgressBar(self.tab_qr, width=400)
        self.progress_bar.set(0)
        self.progress_bar.pack(pady=(5, 0))
        self.progress_bar.pack_forget()

        self.lbl_progress = ctk.CTkLabel(self.tab_qr, text="")
        self.lbl_progress.pack(pady=(0, 5))
        self.lbl_progress.pack_forget()

        self.textbox_qr = ctk.CTkTextbox(self.tab_qr, width=750, height=270, fg_color="#1e1e1e")
        self.textbox_qr.pack(padx=10, pady=10, fill="both", expand=True)

    def pick_excel(self):
        filepath = filedialog.askopenfilename(title="Выберите Excel файл", filetypes=[("Excel", "*.xlsx *.xls")])
        if filepath:
            self.entry_excel.delete(0, 'end')
            self.entry_excel.insert(0, filepath)

    def pick_pdf_in(self):
        dirpath = filedialog.askdirectory(title="Выберите папку с исходными PDF")
        if dirpath:
            self.entry_pdf_in.delete(0, 'end')
            self.entry_pdf_in.insert(0, dirpath)

    def pick_pdf_out(self):
        dirpath = filedialog.askdirectory(title="Выберите папку для сохранения результатов")
        if dirpath:
            self.entry_pdf_out.delete(0, 'end')
            self.entry_pdf_out.insert(0, dirpath)

    def safe_log_qr(self, msg):
        self.after(0, lambda: self._append_log_qr(msg))
        
    def _append_log_qr(self, msg):
        self.textbox_qr.insert("end", msg + "\n")
        self.textbox_qr.see("end")

    def safe_stage_qr(self, msg):
        self.after(0, lambda: self.lbl_stage.configure(text=f"Стадия: {msg}"))

    def safe_progress_qr(self, current, total, filename):
        def update():
            self.progress_bar.set(current / total if total > 0 else 0)
            self.lbl_progress.configure(text=f"{current} / {total} файлов")
            self.lbl_stage.configure(text=f"Проверяю: {filename}")
        self.after(0, update)

    def start_qr_processing(self):
        excel_path = self.entry_excel.get()
        pdf_in = self.entry_pdf_in.get()
        pdf_out = self.entry_pdf_out.get()

        if not excel_path or not pdf_in or not pdf_out:
            self.safe_stage_qr("Ошибка: Заполните все поля путей!")
            return

        self.btn_start.configure(state="disabled")
        self.progress_bar.pack(pady=(5, 0))
        self.lbl_progress.pack(pady=(0, 5))
        self.progress_bar.set(0)
        self.textbox_qr.delete("1.0", "end")

        threading.Thread(target=self.run_qr_script, args=(excel_path, pdf_in, pdf_out), daemon=True).start()

    def run_qr_script(self, current_excel, current_pdf_dir, current_output_dir):
        self.safe_stage_qr("Чтение Excel...")
        self.safe_log_qr("Начинаем сверку QR-кодов...")
        try:
            qrs_to_find = set()
            xls = pd.read_excel(current_excel, sheet_name=None, header=None)
            pattern = re.compile(r'(01\d{14}21.{13})')
            
            for sheet_name, df in xls.items():
                for _, row in df.iterrows():
                    for val in row.dropna():
                        match = pattern.search(str(val).strip())
                        if match:
                            qrs_to_find.add(match.group(1))

            self.safe_log_qr(f"Найдено уникальных кодов в Excel: {len(qrs_to_find)}\n")
            found_qrs = set()
            
            if qrs_to_find:
                pdf_files = glob.glob(os.path.join(current_pdf_dir, '*.pdf'))
                if not pdf_files:
                    self.safe_log_qr("В папке нет PDF файлов.")
                else:
                    total_pages, matched_files, found_qrs = process_all_pdfs(
                        current_pdf_dir, current_output_dir, qrs_to_find, self.safe_log_qr, self.safe_progress_qr
                    )
                    self.safe_log_qr(f"\nСохранено подходящих страниц в PDF: {total_pages}")

            # Генерация отчета по неподтвержденным позициям
            self.safe_stage_qr("Формирую отчет по ненайденным...")
            wb_orig = openpyxl.load_workbook(current_excel)
            wb_new = openpyxl.Workbook()
            wb_new.remove(wb_new.active)
            missing_total = 0

            for sheet_name in wb_orig.sheetnames:
                ws_orig = wb_orig[sheet_name]
                ws_new = wb_new.create_sheet(title=sheet_name)
                for row in ws_orig.iter_rows(values_only=False):
                    row_text = " ".join([str(cell.value).strip() for cell in row if cell.value is not None])
                    if row_text:
                        matches = pattern.findall(row_text)
                        if matches:
                            if not any(m in found_qrs for m in matches):
                                missing_total += 1
                                ws_new.append([cell.value for cell in row])
                        elif row[0].row == 1:
                            ws_new.append([cell.value for cell in row])

            report_path = os.path.join(current_output_dir, "Verification_Report.xlsx")
            wb_new.save(report_path)
            self.safe_log_qr(f"Отчет сохранен: Verification_Report.xlsx (Потеряшек: {missing_total})")
            self.safe_stage_qr("Готово!")

        except Exception as e:
            self.safe_log_qr(f"Ошибка: {e}")
            self.safe_stage_qr("Произошла ошибка")
        finally:
            self.after(0, lambda: self.btn_start.configure(state="normal"))
            self.after(0, self.progress_bar.pack_forget)
            self.after(0, self.lbl_progress.pack_forget)


    # --- ИНТЕРФЕЙС ВКЛАДКИ 2 (ПАРСЕР OZON) ---
    def init_ozon_tab(self):
        header = ctk.CTkLabel(self.tab_ozon, text="Парсер накладных Ozon в Excel", font=ctk.CTkFont(size=22, weight="bold"), text_color="#90CAF9")
        header.pack(pady=(10, 15))

        ozon_frame = ctk.CTkFrame(self.tab_ozon, fg_color="transparent")
        ozon_frame.pack(fill="x", padx=10, pady=5)

        self.btn_ozon_pdf = ctk.CTkButton(ozon_frame, text="Выбрать PDF", width=150, command=self.pick_ozon_pdf)
        self.btn_ozon_pdf.grid(row=0, column=0, padx=(0, 10), pady=10)
        self.entry_ozon_pdf = ctk.CTkEntry(ozon_frame, placeholder_text="PDF файл с накладной Ozon не выбран", width=570)
        self.entry_ozon_pdf.grid(row=0, column=1, pady=10)

        self.btn_ozon_out = ctk.CTkButton(ozon_frame, text="Куда сохранить?", width=150, command=self.pick_ozon_out)
        self.btn_ozon_out.grid(row=1, column=0, padx=(0, 10), pady=10)
        self.entry_ozon_out = ctk.CTkEntry(ozon_frame, placeholder_text="Папка для сохранения результата не выбрана", width=570)
        self.entry_ozon_out.grid(row=1, column=1, pady=10)

        self.btn_start_ozon = ctk.CTkButton(self.tab_ozon, text="Запустить парсинг PDF", height=42, width=250, font=ctk.CTkFont(size=15), command=self.start_ozon_processing)
        self.btn_start_ozon.pack(pady=15)

        self.lbl_stage_ozon = ctk.CTkLabel(self.tab_ozon, text="Стадия: Готово к работе", text_color="#ffcc00", font=ctk.CTkFont(weight="bold"))
        self.lbl_stage_ozon.pack(pady=5)

        self.textbox_ozon = ctk.CTkTextbox(self.tab_ozon, width=750, height=310, fg_color="#1e1e1e")
        self.textbox_ozon.pack(padx=10, pady=10, fill="both", expand=True)

    def pick_ozon_pdf(self):
        filepath = filedialog.askopenfilename(title="Выберите PDF файл Ozon", filetypes=[("PDF", "*.pdf")])
        if filepath:
            self.entry_ozon_pdf.delete(0, 'end')
            self.entry_ozon_pdf.insert(0, filepath)

    def pick_ozon_out(self):
        dirpath = filedialog.askdirectory(title="Выберите папку для сохранения отчета")
        if dirpath:
            self.entry_ozon_out.delete(0, 'end')
            self.entry_ozon_out.insert(0, dirpath)

    def safe_log_ozon(self, msg):
        self.after(0, lambda: self._append_log_ozon(msg))
        
    def _append_log_ozon(self, msg):
        self.textbox_ozon.insert("end", msg + "\n")
        self.textbox_ozon.see("end")

    def safe_stage_ozon(self, msg):
        self.after(0, lambda: self.lbl_stage_ozon.configure(text=f"Стадия: {msg}"))

    def start_ozon_processing(self):
        pdf_path = self.entry_ozon_pdf.get()
        out_dir = self.entry_ozon_out.get()

        if not pdf_path or not out_dir:
            self.safe_stage_ozon("Ошибка: Выберите PDF-файл и папку сохранения!")
            return

        self.btn_start_ozon.configure(state="disabled")
        self.textbox_ozon.delete("1.0", "end")

        threading.Thread(target=self.run_ozon_script, args=(pdf_path, out_dir), daemon=True).start()

    def run_ozon_script(self, pdf_path, out_dir):
        self.safe_stage_ozon("Парсинг PDF...")
        self.safe_log_ozon(f"Читаем файл: {os.path.basename(pdf_path)}")
        try:
            df_orders = parse_orders_from_pdf(pdf_path)
            if df_orders.empty:
                self.safe_log_ozon("Не удалось найти заказы или товары по заданным шаблонам.")
                self.safe_stage_ozon("Завершено с предупреждением")
            else:
                output_xlsx = os.path.join(out_dir, "Ozon_Orders_List.xlsx")
                save_to_formatted_excel(df_orders, output_xlsx)
                self.safe_log_ozon(f"Успешно обработано позиций: {len(df_orders)}")
                self.safe_log_ozon(f"Файл сохранен: {output_xlsx}")
                self.safe_stage_ozon("Успешно завершено!")
        except Exception as e:
            self.safe_log_ozon(f"Ошибка при обработке PDF: {e}")
            self.safe_stage_ozon("Произошла ошибка")
        finally:
            self.after(0, lambda: self.btn_start_ozon.configure(state="normal"))


if __name__ == "__main__":
    app = MainApp()
    app.mainloop()
